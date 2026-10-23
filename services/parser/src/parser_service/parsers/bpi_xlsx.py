"""BPI Net account statement (xlsx, "Extracto movimentos").

Sheet `Movimentos`: a few title and holder rows, the table (`Data Mov.`, `Data Valor`, `Descrição do
Movimento`, `Valor em EUR`, `Saldo em EUR`) and a note under it. Dates are `dd-mm-aaaa` text, the
amount carries its sign and the file lists the newest movement first; rows come out oldest first so
each balance follows the one before it. The movement date is the row date; the value date is only
informative (card purchases state their own date at the start of the description).
"""

import io
import itertools
import re
import warnings
from decimal import Decimal
from typing import Any

import openpyxl

from ..errors import ParseFailedError, UnsupportedFileError
from ..helpers import normalize_text, parse_date, parse_num
from ..models import BankRow, BankRowType, ParseMeta, ParseResponse, ParserKind
from ..sniff import SniffedFile

SHEET = "Movimentos"
_DATE_COLUMN = "DATA MOV."
_DESCRIPTION_COLUMN = "DESCRICAO DO MOVIMENTO"
_AMOUNT_COLUMN = re.compile(r"^VALOR EM ([A-Z]{3})$")
_BALANCE_COLUMN = re.compile(r"^SALDO EM [A-Z]{3}$")
_ACCOUNT = re.compile(r"^\s*([0-9][0-9A-Za-z.\-]*)\s*\(([A-Z]{3})\)")
_TYPES: tuple[tuple[re.Pattern[str], BankRowType], ...] = (
    (re.compile(r"^(?:\d{2}/\d{2} )?COMPRA ELEC\b"), "card"),
    (re.compile(r"^TRF\b"), "transfer"),
    (re.compile(r"^(?:IMPOSTO DE SELO|MANUTENCAO DE CONTA)\b"), "fee"),
)


class BpiXlsx:
    id = "bpi-xlsx"
    kind: ParserKind = "bank"
    version = "1.0.0"

    def detect(self, file: SniffedFile) -> float:
        if file.magic != "xlsx" or SHEET not in file.sheet_names:
            return 0.0
        head = file.header_rows.get(SHEET, [])
        has_table = _header_at(head) is not None
        branded = any(row and row[0] == "BPI Net" for row in head)
        if has_table:
            return 0.95 if branded else 0.75
        return 0.5 if branded else 0.0

    def parse(self, file: SniffedFile) -> ParseResponse:
        workbook = _open(file.data)
        if SHEET not in workbook.sheetnames:
            raise ParseFailedError('The "Movimentos" sheet was not found.')
        sheet = [list(r) for r in workbook[SHEET].iter_rows(values_only=True)]
        header_at = _header_at(sheet)
        if header_at is None:
            raise ParseFailedError("The movements table was not found.")
        header = [normalize_text(str(c)) if c is not None else "" for c in sheet[header_at]]
        columns = _columns(header)
        currency = _AMOUNT_COLUMN.match(header[columns["amount"]])
        warnings_out: list[str] = []
        if currency and currency[1] != "EUR":
            warnings_out.append(f"The statement is in {currency[1]}: confirm the account currency.")

        listed: list[BankRow] = []
        ignored = 0
        for raw in sheet[header_at + 1 :]:
            day = parse_date(_text(_at(raw, columns["date"])))
            if day is None:
                continue  # blank line and the note under the table
            amount = parse_num(_at(raw, columns["amount"]))
            description = " ".join(_text(_at(raw, columns["description"])).split())
            if amount is None or not description:
                ignored += 1
                continue
            balance = parse_num(_at(raw, columns["balance"]))
            listed.append(
                BankRow(
                    kind="bank",
                    date=day,
                    description=description,
                    amount=amount,
                    balance=balance,
                    type=_type(description),
                )
            )
        if ignored:
            warnings_out.append(f"{ignored} rows without an amount or description were skipped.")
        if not listed:
            raise ParseFailedError("No movements were found in this statement.")

        rows = list(reversed(listed))  # newest first in the file
        reconciled = _reconciled(rows)
        if reconciled is False or ignored:
            reconciled = False
            warnings_out.append("The running balances do not add up.")
        first, last = rows[0], rows[-1]
        opening = (
            None
            if first.balance is None
            else format(Decimal(first.balance) - Decimal(first.amount), "f")
        )
        return ParseResponse(
            rows=list(rows),
            meta=ParseMeta(
                parser_id=self.id,
                parser_version=self.version,
                row_count=len(rows),
                opening_balance=opening and _plain(opening),
                closing_balance=last.balance,
                reconciled=reconciled,
                account_hint=_account(sheet[:header_at]),
            ),
            warnings=warnings_out,
        )


def _open(data: bytes) -> openpyxl.Workbook:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise UnsupportedFileError("The workbook could not be opened.") from exc


def _header_at(rows: list[list[Any]]) -> int | None:
    for index, row in enumerate(rows):
        names = [normalize_text(str(c)) for c in row if c is not None]
        if (
            _DATE_COLUMN in names
            and _DESCRIPTION_COLUMN in names
            and any(_AMOUNT_COLUMN.match(n) for n in names)
            and any(_BALANCE_COLUMN.match(n) for n in names)
        ):
            return index
    return None


def _columns(header: list[str]) -> dict[str, int]:
    found = {
        "date": _DATE_COLUMN,
        "description": _DESCRIPTION_COLUMN,
    }
    columns = {key: header.index(name) for key, name in found.items()}
    columns["amount"] = next(i for i, n in enumerate(header) if _AMOUNT_COLUMN.match(n))
    columns["balance"] = next(i for i, n in enumerate(header) if _BALANCE_COLUMN.match(n))
    return columns


def _reconciled(rows: list[BankRow]) -> bool | None:
    """Each balance equals the one before plus the amount; `None` when no balance is stated."""
    checked = False
    for before, row in itertools.pairwise(rows):
        if before.balance is None or row.balance is None:
            continue
        checked = True
        if Decimal(before.balance) + Decimal(row.amount) != Decimal(row.balance):
            return False
    return True if checked else None


def _account(title_rows: list[list[Any]]) -> str | None:
    for row in title_rows:
        if row[:1] and normalize_text(_text(row[0])) == "CONTA E MOEDA":
            match = next(
                (m for c in row[1:] if (m := _ACCOUNT.match(_text(c)))),
                None,
            )
            return match[1] if match else None
    return None


def _type(description: str) -> BankRowType | None:
    key = normalize_text(description)
    return next((kind for pattern, kind in _TYPES if pattern.match(key)), None)


def _at(row: list[Any], index: int) -> Any:
    return row[index] if index < len(row) else None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _plain(value: str) -> str:
    return format(Decimal(value).normalize(), "f") if Decimal(value) != 0 else "0"
