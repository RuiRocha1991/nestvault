"""`SniffedFile`: what every parser needs to recognise a file, computed once per request."""

import csv
import io
import zipfile
from dataclasses import dataclass, field
from typing import Literal

import openpyxl
import pdfplumber

from .config import Settings
from .errors import UnsupportedFileError
from .limits import check_zip

MagicType = Literal["xlsx", "pdf", "csv"]
Cell = str | int | float | None

MAX_SHEETS = 20
HEADER_ROWS = 15
MAX_COLUMNS = 60
PDF_PAGES = 2
PDF_LINES = 60
CSV_ENCODINGS = ("utf-8-sig", "cp1252")


@dataclass(frozen=True)
class SniffedFile:
    data: bytes
    magic: MagicType
    sheet_names: tuple[str, ...] = ()
    # First rows per sheet (xlsx) or of the file under the key "" (csv). Dates become ISO text.
    header_rows: dict[str, list[list[Cell]]] = field(default_factory=dict)
    pdf_lines: tuple[str, ...] = ()
    text: str | None = None  # decoded csv content


def sniff(data: bytes, settings: Settings) -> SniffedFile:
    if data.startswith(b"%PDF-"):
        return _sniff_pdf(data)
    if data.startswith(b"PK\x03\x04"):
        return _sniff_xlsx(data, settings)
    return _sniff_csv(data)


def _cell(value: object) -> Cell:
    if value is None or isinstance(value, str | int | float):
        return value
    isoformat = getattr(value, "isoformat", None)
    return str(isoformat()) if callable(isoformat) else str(value)


def _sniff_xlsx(data: bytes, settings: Settings) -> SniffedFile:
    check_zip(data, settings)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = set(archive.namelist())
    if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
        raise UnsupportedFileError("Only Excel (.xlsx) workbooks are supported.")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise UnsupportedFileError("The workbook could not be opened.") from exc
    try:
        rows: dict[str, list[list[Cell]]] = {}
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            head: list[list[Cell]] = []
            for row in sheet.iter_rows(max_row=HEADER_ROWS, max_col=MAX_COLUMNS, values_only=True):
                head.append([_cell(value) for value in row])
            rows[sheet.title] = head
        return SniffedFile(
            data=data,
            magic="xlsx",
            sheet_names=tuple(workbook.sheetnames[:MAX_SHEETS]),
            header_rows=rows,
        )
    finally:
        workbook.close()


def _sniff_pdf(data: bytes) -> SniffedFile:
    lines: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:PDF_PAGES]:
                lines.extend((page.extract_text() or "").splitlines())
    except Exception as exc:
        raise UnsupportedFileError("The PDF could not be read.") from exc
    return SniffedFile(data=data, magic="pdf", pdf_lines=tuple(lines[:PDF_LINES]))


def _sniff_csv(data: bytes) -> SniffedFile:
    if b"\x00" in data[:4096]:
        raise UnsupportedFileError()
    text: str | None = None
    for encoding in CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None or not text.strip():
        raise UnsupportedFileError()
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    head: list[list[Cell]] = []
    for row in reader:
        head.append(list(row[:MAX_COLUMNS]))
        if len(head) >= HEADER_ROWS:
            break
    return SniffedFile(data=data, magic="csv", header_rows={"": head}, text=text)
