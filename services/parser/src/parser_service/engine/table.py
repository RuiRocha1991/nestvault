"""Rows out of a csv/xlsx file, finding the header row, and guessing which column is which
(ports of `parseCSV`, `findHeader` and `guessMap` in the prototype)."""

import csv
import io
import re
import warnings
from typing import Any

import openpyxl

from ..errors import ParseFailedError, UnsupportedFileError
from ..helpers import normalize_text, parse_date
from ..sniff import SniffedFile

MAX_ROWS = 200_000
SCAN_ROWS = 40
DELIMITERS = (";", ",", "\t", "|")
Rows = list[list[Any]]
_DATE_WORDS = re.compile(r"DATA|DATE")
_DESCRIPTION_WORDS = re.compile(r"DESCRI|MOVIMENTO|HISTORICO|CONCEITO|DESCRIPTION|DETALHE")


def read_rows(
    file: SniffedFile, *, sheet: str | None, delimiter: str | None, encoding: str | None
) -> Rows:
    if file.magic == "xlsx":
        return _xlsx_rows(file.data, sheet)
    if file.magic == "csv":
        return _csv_rows(file, delimiter, encoding)
    raise UnsupportedFileError("This spec reads csv and xlsx files.")


def _xlsx_rows(data: bytes, sheet: str | None) -> Rows:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise UnsupportedFileError("The workbook could not be opened.") from exc
    if sheet is None:
        worksheet = workbook.worksheets[0]
    elif sheet in workbook.sheetnames:
        worksheet = workbook[sheet]
    else:
        raise ParseFailedError("The sheet this spec reads is not in the workbook.")
    rows: Rows = []
    for row in worksheet.iter_rows(values_only=True):
        rows.append([c.strip() if isinstance(c, str) else c for c in row])
        if len(rows) > MAX_ROWS:
            raise ParseFailedError("The sheet has too many rows.")
    return rows


def detect_delimiter(text: str) -> str:
    lines = text.replace("\r", "").split("\n")[:SCAN_ROWS]
    return max(DELIMITERS, key=lambda d: sum(line.count(d) for line in lines))


def _csv_rows(file: SniffedFile, delimiter: str | None, encoding: str | None) -> Rows:
    text = file.data.decode(encoding, errors="replace") if encoding else file.text
    if text is None:
        raise UnsupportedFileError()
    reader = csv.reader(io.StringIO(text), delimiter=delimiter or detect_delimiter(text))
    rows: Rows = []
    for row in reader:
        rows.append([cell.strip() for cell in row])
        if len(rows) > MAX_ROWS:
            raise ParseFailedError("The file has too many rows.")
    return rows


def is_blank(row: list[Any]) -> bool:
    return all(cell in (None, "") for cell in row)


def find_header(rows: Rows) -> int:
    """Index of the header row: the first row that looks like `date ... description`, else the row
    just before the first dated row."""
    window = rows[:SCAN_ROWS]
    for index, row in enumerate(window):
        joined = "|".join(normalize_text(str(c)) for c in row if c is not None)
        if _DATE_WORDS.search(joined) and _DESCRIPTION_WORDS.search(joined):
            return index
    for index in range(len(window) - 1):
        if _is_date(window[index + 1][0] if window[index + 1] else None) and not _is_date(
            window[index][0] if window[index] else None
        ):
            return index
    return 0


def _is_date(cell: object) -> bool:
    return isinstance(cell, str | int | float) and parse_date(cell) is not None


def header_matching(rows: Rows, headers: list[str]) -> int | None:
    """First row that has every one of `headers` (compared without accents or case)."""
    wanted = {normalize_text(h) for h in headers}
    for index, row in enumerate(rows[:SCAN_ROWS]):
        if wanted <= {normalize_text(str(c)) for c in row if c is not None}:
            return index
    return None


def column_index(header: list[Any], name: str) -> int | None:
    key = normalize_text(name)
    for index, cell in enumerate(header):
        if cell is not None and normalize_text(str(cell)) == key:
            return index
    return None


def guess_columns(header: list[Any]) -> dict[str, int | None]:
    """Which header cell is the date, description, amount, debit, credit and balance."""
    cells = [normalize_text(str(c)) if c is not None else "" for c in header]

    def find(*patterns: str) -> int | None:
        for pattern in patterns:
            for index, cell in enumerate(cells):
                if re.search(pattern, cell):
                    return index
        return None

    return {
        "date": find(r"DATA (MOV|LANC|OPER)", r"^DATA$", r"DATA", r"DATE"),
        "description": find(
            r"DESCRI", r"MOVIMENTO", r"HISTORICO", r"DETALHE", r"CONCEITO", r"DESCRIPTION"
        ),
        "amount": find(r"^(MONTANTE|VALOR|IMPORTANCIA|AMOUNT)", r"MONTANTE", r"^VALOR(?!.*SALDO)"),
        "debit": find(r"DEBITO", r"DEBIT"),
        "credit": find(r"CREDITO", r"CREDIT"),
        "balance": find(r"^SALDO", r"BALANCE"),
    }
